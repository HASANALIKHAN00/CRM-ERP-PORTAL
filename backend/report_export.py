import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def flatten_records(records: list) -> list:
    """
    Shared across Excel/CSV/PDF export: flattens any nested 'detail' lists
    (e.g. get_team_performance_summary) into a readable flat row shape.
    Single source of truth so all three export formats stay consistent.
    """
    if not records:
        return []
    flat = []
    for rec in records:
        if isinstance(rec, dict) and "detail" in rec and isinstance(rec["detail"], list):
            for sub in rec["detail"]:
                flat.append({"metric": rec.get("metric"), **sub})
        elif isinstance(rec, dict) and "detail" in rec:
            flat.append({"metric": rec.get("metric"), "value": rec.get("detail")})
        else:
            flat.append(rec)
    return flat


def build_excel(result: dict, function_name: str) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = function_name[:31]

    bold = Font(bold=True)
    header_fill = PatternFill(start_color="0F241A", end_color="0F241A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    ws["A1"] = "Summary"
    ws["A1"].font = bold
    ws["A2"] = result.get("summary", "")
    ws.merge_cells("A2:F2")
    ws["A2"].alignment = ws["A2"].alignment.copy(wrap_text=True)

    ws["A4"] = "Source"
    ws["A4"].font = bold
    ws["A5"] = result.get("source", "")
    ws.merge_cells("A5:F5")

    ws["A6"] = "Generated"
    ws["A6"].font = bold
    ws["B6"] = datetime.now().strftime("%Y-%m-%d %H:%M")

    flat_records = flatten_records(result.get("records") or [])
    start_row = 8

    if not flat_records:
        ws.cell(row=start_row, column=1, value="No records available for this report/filter combination.")
    else:
        headers = list(flat_records[0].keys())
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header.replace("_", " ").title())
            cell.font = header_font
            cell.fill = header_fill

        for row_idx, rec in enumerate(flat_records, start=start_row + 1):
            for col_idx, header in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=col_idx, value=rec.get(header))

        for col_idx, header in enumerate(headers, start=1):
            max_len = max([len(str(header))] + [len(str(r.get(header, ""))) for r in flat_records])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_csv(result: dict) -> io.BytesIO:
    flat_records = flatten_records(result.get("records") or [])

    text_buffer = io.StringIO()
    text_buffer.write(f"# Summary: {result.get('summary', '')}\n")
    text_buffer.write(f"# Source: {result.get('source', '')}\n")
    text_buffer.write(f"# Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n")

    if not flat_records:
        text_buffer.write("No records available for this report/filter combination.\n")
    else:
        headers = list(flat_records[0].keys())
        writer = csv.DictWriter(text_buffer, fieldnames=headers)
        writer.writeheader()
        for rec in flat_records:
            writer.writerow({h: rec.get(h, "") for h in headers})

    byte_buffer = io.BytesIO(text_buffer.getvalue().encode("utf-8"))
    byte_buffer.seek(0)
    return byte_buffer


def build_pdf(result: dict, function_name: str) -> io.BytesIO:
    buffer = io.BytesIO()
    flat_records = flatten_records(result.get("records") or [])

    # Landscape for wide tables (many columns), portrait for narrow ones --
    # a report with 2-3 columns doesn't need the extra width, and a report
    # with 6+ columns needs it to stay readable without tiny text.
    page_size = landscape(letter) if flat_records and len(flat_records[0]) > 4 else letter
    doc = SimpleDocTemplate(buffer, pagesize=page_size, topMargin=0.6 * inch, bottomMargin=0.6 * inch)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=16, spaceAfter=6)
    meta_style = ParagraphStyle("ReportMeta", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))
    summary_style = ParagraphStyle("ReportSummary", parent=styles["Normal"], fontSize=11, spaceAfter=4)

    story = []
    story.append(Paragraph(function_name.replace("_", " ").title(), title_style))
    story.append(Paragraph(result.get("summary", ""), summary_style))
    story.append(Paragraph(f"Source: {result.get('source', '')}", meta_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", meta_style))
    story.append(Spacer(1, 16))

    if not flat_records:
        story.append(Paragraph("No records available for this report/filter combination.", styles["Normal"]))
    else:
        headers = list(flat_records[0].keys())
        header_row = [h.replace("_", " ").title() for h in headers]
        table_data = [header_row]
        for rec in flat_records:
            table_data.append([str(rec.get(h, "")) for h in headers])

        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F241A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f2f2")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)

    doc.build(story)
    buffer.seek(0)
    return buffer
