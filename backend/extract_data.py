import openpyxl
import json
from datetime import datetime

wb = openpyxl.load_workbook(r'C:\Users\SUN RISE\OneDrive\Desktop\BoxTech_Activity_Graph_UI_With_Top_Cards.xlsx', data_only=True)
s = wb['Activity Dashboard']

records = []
for r in range(2, s.max_row + 1):
    dt_val = s.cell(r, 24).value
    salesperson = s.cell(r, 25).value
    country = s.cell(r, 26).value
    activity = s.cell(r, 27).value
    client = s.cell(r, 28).value
    date_only = s.cell(r, 29).value
    hour_slot = s.cell(r, 30).value
    half_hour_slot = s.cell(r, 31).value
    amount = s.cell(r, 36).value
    
    if dt_val is None and salesperson is None:
        continue
        
    if isinstance(dt_val, datetime):
        dt_str = dt_val.strftime('%Y-%m-%d %H:%M:%S')
        d_str = dt_val.strftime('%Y-%m-%d')
        t_str = dt_val.strftime('%H:%M')
    else:
        dt_str = str(dt_val)
        d_str = dt_str[:10]
        t_str = dt_str[11:16] if len(dt_str) > 11 else '00:00'
        
    h_slot = int(hour_slot) if hour_slot is not None else int(t_str.split(':')[0])
    hh_slot = int(half_hour_slot) if half_hour_slot is not None else (h_slot * 2 + (1 if int(t_str.split(':')[1]) >= 30 else 0))
    
    m = int(t_str.split(':')[1]) if ':' in t_str else 0
    hh_str = f'{h_slot:02d}:30' if m >= 30 else f'{h_slot:02d}:00'
    
    amt_f = float(amount) if amount is not None else 0.0
    
    records.append({
        'id': len(records) + 1,
        'date': d_str,
        'time': t_str,
        'hourSlot': h_slot,
        'halfHourSlot': hh_str,
        'salesperson': str(salesperson).strip() if salesperson else '',
        'country': str(country).strip() if country else '',
        'activityType': str(activity).strip() if activity else '',
        'client': str(client).strip() if client else '',
        'amount': amt_f,
        'isFinancial': amt_f > 0 or str(activity).strip() in {'Quotation', 'Sales Order', 'Sales Invoice', 'Payment Entry', 'Sales Opportunity'}
    })

print(f'Total records extracted: {len(records)}')

p1 = [r for r in records if r['date'] == '2026-08-08' and r['activityType'] == 'Phone Call']
print(f'2026-08-08 Phone Call: count={len(p1)}, clients={len(set(r["client"] for r in p1))}')

with open(r'C:\Users\SUN RISE\OneDrive\Desktop\my_project\src\app\data\actual-activity-records.json', 'w') as f:
    json.dump(records, f, indent=2)

print('Saved to src/app/data/actual-activity-records.json successfully')
